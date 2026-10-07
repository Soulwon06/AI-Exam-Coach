"""Course Brain: Material → Evidence → Knowledge Domains → Study Units → Questions →
Question Mapping → Scoring. See docs/04_COURSE_BRAIN.md and docs/08_AI_PIPELINE.md.
"""
import os
import re
import hashlib
import threading

from sqlalchemy.orm import Session

from app import models
from app.core.errors import AIProcessingFailedError
from app.services import ai_service
from app.services.exam_intelligence import process_exam_material
from app.services.knowledge_base import TEACHING_TYPES

PARSER_VERSION = "material-parser-v1"
_PARSER_CACHE_STATUSES = {"PARSED", "LOW_QUALITY"}


def process_material(db: Session, material: models.Material) -> models.Material:
    """Run the AI pipeline over one material and persist results. Called by the
    background worker. Never throws for optional-source failure (graceful degradation).
    """
    try:
        parsed = parse_material(db, material)
        text = parsed["text"]
        if parsed["status"] in {"FAILED", "UNSUPPORTED"}:
            material.status = "FAILED"
            material.progress = 100
            material.error_code = parsed["error"] or "PARSER_FAILED"
            db.flush()
            return material
        # REAL_EXAM is question evidence, not a Knowledge Base source.  Keep
        # the parser/cache step above for newly uploaded files, then route the
        # persisted text through the deterministic exam pipeline.  That path
        # never creates StudyUnits and never asks the AI provider to invent one.
        if material.type == "REAL_EXAM":
            exam_result = process_exam_material(db, material)
            if exam_result["status"] == "FAILED":
                material.status = "FAILED"
                material.progress = 100
                material.error_code = "EXAM_NO_CACHED_TEXT"
            else:
                material.status = "COMPLETED"
                material.progress = 100
                material.error_code = None
            db.flush()
            return material
        # Teaching StudyUnits are generated once, at the course-level finalize
        # boundary, by the canonical cache-only Knowledge Base builder.  Do
        # not invoke the legacy AI provider or create per-material units here.
        if material.type in TEACHING_TYPES:
            material.status = "COMPLETED"
            material.progress = 100
            material.error_code = None
            db.flush()
            return material
        # Scanned/unreadable material: no analyzable text -> skip the analyzer
        # entirely so it never fabricates units, questions, or a fake evidence
        # from the degradation message itself.
        if parsed["status"] == "PARSED" and text.strip():
            result = ai_service.analyze_material(text, material.type)
        else:
            result = {"study_units": [], "questions": [], "evidence": []}
        material.status = "COMPLETED"
        material.progress = 100
        material.error_code = None

        course_id = material.course_id
        # Do not attach one material-wide summary to every unit.  That was the
        # original source of false "evidence": any line extracted from a PDF
        # appeared to be supported by the whole paper.  Each unit is now linked
        # only to its own quoted source snippet.
        _create_evidence(db, material, result.get("evidence", []))
        units = _create_study_units(db, course_id, result.get("study_units", []), material)
        questions = _create_questions(db, course_id, result.get("questions", []))
        _map_questions(db, course_id, questions, units, material)

        db.flush()
        return material
    except Exception as exc:  # graceful degradation
        material.status = "FAILED"
        material.error_code = "AI_PROCESSING_FAILED"
        db.flush()
        raise AIProcessingFailedError(f"Material processing failed: {exc}")


def _file_sha256(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _quality_for_text(text: str) -> tuple[str, str, str | None]:
    text = _sanitize_text(text or "")
    chars = len(text)
    if chars == 0:
        return "FAILED", "NONE", "no usable text extracted"
    ratio = _readable_ratio(text)
    if chars < 20 or ratio < 0.5:
        return "LOW_QUALITY", "LOW", f"extracted text is too short or unreadable (chars={chars}, readable_ratio={ratio:.2f})"
    if chars < 100 or ratio < 0.75:
        return "LOW_QUALITY", "LOW", f"extracted text is below normal quality (chars={chars}, readable_ratio={ratio:.2f})"
    return "PARSED", "HIGH", None


def _set_parse_result(material: models.Material, *, file_hash: str | None,
                      method: str, text: str, status: str, quality: str,
                      error: str | None) -> dict:
    clean = _sanitize_text(text or "")
    material.file_hash = file_hash
    material.parser_version = PARSER_VERSION
    material.parser_method = method
    material.extracted_text = clean
    material.extracted_char_count = len(clean)
    material.extraction_quality = quality
    material.extraction_error = error
    material.parsing_status = status
    return {"text": clean, "status": status, "quality": quality,
            "method": method, "error": error, "cached": False}


def parse_material(db: Session, material: models.Material, force: bool = False) -> dict:
    """Extract and persist source text, reusing a stable successful cache.

    This function intentionally stops at parsing.  Knowledge-point generation
    remains in process_material and is not called by the P1-B audit script.
    """
    previous_parsing_status = material.parsing_status
    material.parsing_status = "PROCESSING"
    db.flush()

    if material.external_url:
        return _set_parse_result(
            material, file_hash=None, method="external_url", text="",
            status="FAILED", quality="NONE",
            error="external URL has no local text parser",
        )
    if not material.file_path:
        return _set_parse_result(
            material, file_hash=None, method="none", text="",
            status="FAILED", quality="NONE", error="material has no local file",
        )

    path = _resolve_material_path(material.file_path)
    if not os.path.isfile(path):
        return _set_parse_result(
            material, file_hash=None, method="file", text="",
            status="FAILED", quality="NONE", error=f"file not found: {material.name}",
        )
    if path != material.file_path:
        material.file_path = path

    file_hash = _file_sha256(path)
    cached_status = previous_parsing_status
    # A prior process could have died after persisting the text fields but
    # before restoring PROCESSING.  Recover only when hash/version and a
    # non-empty persisted result make the status unambiguous.
    if (
        previous_parsing_status == "PROCESSING"
        and material.extracted_text
        and material.extracted_char_count
        and material.extraction_quality in {"HIGH", "LOW"}
    ):
        cached_status = "PARSED" if material.extraction_quality == "HIGH" else "LOW_QUALITY"
    if (
        not force
        and material.file_hash == file_hash
        and material.parser_version == PARSER_VERSION
        and cached_status in _PARSER_CACHE_STATUSES
        and material.extracted_text is not None
    ):
        material.parsing_status = cached_status
        return {
            "text": material.extracted_text,
            "status": cached_status,
            "quality": material.extraction_quality,
            "method": material.parser_method,
            "error": material.extraction_error,
            "cached": True,
        }

    extension = os.path.splitext(path)[1].lower()
    try:
        if extension == ".pdf":
            text, method, parser_error = _extract_pdf_text_with_method(path)
        elif extension == ".docx":
            text, method, parser_error = _extract_docx_text_with_method(path)
        elif extension in {".txt", ".md"}:
            with open(path, "r", encoding="utf-8-sig", errors="replace") as source:
                text = source.read()
            method, parser_error = "text", None
        elif extension == ".doc":
            return _set_parse_result(
                material, file_hash=file_hash, method="unsupported_doc", text="",
                status="UNSUPPORTED", quality="NONE",
                error="legacy .doc (OLE binary) is unsupported; upload .docx or PDF",
            )
        elif extension in {".ppt", ".pptx"}:
            return _set_parse_result(
                material, file_hash=file_hash, method="unsupported_ppt", text="",
                status="UNSUPPORTED", quality="NONE",
                error="PPT/PPTX parser is not implemented in this MVP",
            )
        else:
            return _set_parse_result(
                material, file_hash=file_hash, method="unsupported", text="",
                status="UNSUPPORTED", quality="NONE",
                error=f"unsupported file extension: {extension or '(none)'}",
            )
    except Exception as exc:
        return _set_parse_result(
            material, file_hash=file_hash, method=extension.lstrip(".") or "unknown",
            text="", status="FAILED", quality="NONE",
            error=f"{type(exc).__name__}: {exc}",
        )

    text = _sanitize_text(text or "")
    status, quality, quality_error = _quality_for_text(text)
    error = parser_error or quality_error
    result = _set_parse_result(
        material, file_hash=file_hash, method=method, text=text,
        status=status, quality=quality, error=error,
    )
    db.flush()
    return result


# Characters that count as readable text for the scan-vs-binary heuristic.
_READABLE_RE = re.compile(r"[\u4e00-\u9fff\w\u3000-\u303f\uff00-\uffef\s.,;:!?()（）【】《》\-\"'…·、]")
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b-\x1f\x7f]")
# Lone Unicode surrogates (e.g. math chars from pypdf like \ud835) can't be stored
# in SQLite/UTF-8. Drop them before persistence.
_SURROGATE_RE = re.compile(r"[\ud800-\udfff]")


def _sanitize_text(text: str) -> str:
    return _SURROGATE_RE.sub("", text)


def _readable_ratio(text: str) -> float:
    """Fraction of characters that look like real text. Binary/scan fragments
    (compressed JPEG data, PDF object markers) score low."""
    text = _sanitize_text(text)
    if not text:
        return 0.0
    good = sum(1 for ch in text if _READABLE_RE.match(ch))
    return good / len(text)


# --- OCR fallback (scanned PDFs / image-only .docx) -------------------------
# RapidOCR (onnxruntime) is heavy (~a few hundred MB), so load the engine once
# as a lazy singleton instead of once per material. Guard with a lock: the
# worker pool may process several scanned files concurrently.
_ocr_lock = threading.Lock()
_ocr_engine = None


def _get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        with _ocr_lock:
            if _ocr_engine is None:
                from rapidocr_onnxruntime import RapidOCR

                _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_image_bytes(img_bytes: bytes) -> str:
    """OCR a single raster image (PNG/JPEG bytes) into text lines."""
    try:
        import io

        from PIL import Image

        Image.open(io.BytesIO(img_bytes)).convert("RGB")
        engine = _get_ocr_engine()
        result, _ = engine(img_bytes)
        if not result:
            return ""
        lines = []
        for box, txt, conf in result:
            line = (txt or "").strip()
            if line:
                lines.append(line)
        return "\n".join(lines)
    except Exception:
        return ""


def _ocr_pdf(path: str) -> str:
    """Render each page of a scanned PDF to an image and OCR it."""
    try:
        import pymupdf

        doc = pymupdf.open(path)
    except Exception:
        return ""
    parts = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=200)
        text = _ocr_image_bytes(pix.tobytes("png"))
        if text:
            parts.append(f"----- page {i + 1} -----")
            parts.append(text)
    return "\n".join(parts)


def _ocr_docx(path: str) -> str:
    """OCR every embedded image in a scanned .docx (no text layer)."""
    try:
        import zipfile

        with zipfile.ZipFile(path) as zf:
            media = sorted(
                n for n in zf.namelist() if n.lower().startswith("word/media/")
            )
    except Exception:
        return ""
    parts = []
    for name in media:
        try:
            with zipfile.ZipFile(path) as zf:
                data = zf.read(name)
        except Exception:
            continue
        text = _ocr_image_bytes(data)
        if text:
            parts.append(f"----- image {name} -----")
            parts.append(text)
    return "\n".join(parts)


def _read_material_text(material: models.Material) -> str:
    if material.external_url:
        return f"External link material: {material.external_url}\nThe student referenced this as {material.type} evidence."
    if material.file_path:
        path = _resolve_material_path(material.file_path)
        try:
            if path.lower().endswith(".pdf"):
                text = _extract_pdf_text(path)
            elif path.lower().endswith(".docx"):
                text = _extract_docx_text(path)
            elif path.lower().endswith(".doc"):
                text = _extract_doc_text(path)
            else:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
        except FileNotFoundError:
            # Return no text, not an English placeholder that could be mistaken
            # for course evidence. The worker records a precise failure status.
            raise FileNotFoundError(f"Material file unavailable: {material.name}")
        if path != material.file_path:
            # Databases created under WSL/Linux may retain /mnt/d/... paths.
            # The upload basename is immutable, so migrate it to this checkout's
            # data directory once it has been verified to exist.
            material.file_path = path
        return _sanitize_text(text)
    return f"Material: {material.name} (type={material.type})"


def _resolve_material_path(path: str) -> str:
    """Resolve a persisted upload path after a checkout moves OS or location.

    Only the basename is used for fallback, and it must exist inside this
    application's managed upload directory; arbitrary external paths are never
    searched or rewritten.
    """
    if os.path.exists(path):
        return path
    filename = os.path.basename((path or "").replace("\\", "/"))
    if not filename:
        return path
    local_path = os.path.abspath(
        os.path.join(os.path.dirname(__file__), "..", "..", "data", "materials", filename)
    )
    return local_path if os.path.isfile(local_path) else path


def _extract_doc_text(path: str) -> str:
    """Legacy binary .doc (OLE compound). No lightweight pure-Python extractor
    is bundled and page rendering requires external tools; degrade honestly
    instead of feeding raw binary to the analyzer."""
    return ""


def _extract_docx_text_with_method(path: str) -> tuple[str, str, str | None]:
    """Return DOCX text and the actual parser method used."""
    try:
        from docx import Document

        doc = Document(path)
        parts = [par.text.strip() for par in doc.paragraphs if par.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    txt = cell.text.strip()
                    if txt:
                        parts.append(txt)
    except Exception as exc:
        parts = []
        parse_error = f"docx text parser failed: {type(exc).__name__}: {exc}"
    else:
        parse_error = None
    text = "\n".join(parts)
    if len(text.strip()) >= 20 and _readable_ratio(text) >= 0.5:
        return text, "docx", parse_error
    ocr_text = _ocr_docx(path)
    if ocr_text:
        return ocr_text, "ocr", None
    return text, "docx", parse_error or "DOCX text layer is empty and OCR returned no usable text"


def _extract_docx_text(path: str) -> str:
    """Extract paragraphs + table cells from a .docx (OOXML zip). If there is no
    usable text layer (scanned/image-only docx), fall back to OCR of the images."""
    text, _, _ = _extract_docx_text_with_method(path)
    return text if len(text.strip()) >= 20 and _readable_ratio(text) >= 0.5 else ""


def _extract_pdf_text(path: str) -> str:
    """Extract a real text layer with pypdf. Scanned/encrypted PDFs yield little
    or nothing — fall back to OCR (page-by-page image rendering). Raw binary
    fragments (JPEG data, '10 0 obj' markers) must never reach the analyzer,
    which used to fabricate garbage knowledge points."""
    text, _, _ = _extract_pdf_text_with_method(path)
    return text


def _extract_pdf_text_with_method(path: str) -> tuple[str, str, str | None]:
    """Extract PDF text layer first, then OCR only when quality is insufficient."""
    parser_error = None
    try:
        from pypdf import PdfReader

        reader = PdfReader(path)
        pages = []
        text_layer_parts = []
        for index, page in enumerate(reader.pages):
            # Keep page boundaries so the provider can identify repeated
            # headers/footers instead of treating them as course concepts.
            pages.append(f"----- page {index + 1} -----")
            page_text = str(page.extract_text() or "")
            pages.append(page_text)
            text_layer_parts.append(page_text)
    except Exception as exc:
        pages = []
        text_layer_parts = []
        parser_error = f"PDF text parser failed: {type(exc).__name__}: {exc}"
    text = "\n".join(pages)
    text_layer = "\n".join(text_layer_parts)
    # Page markers are control metadata, not evidence. Exclude them from the
    # text-layer quality check or a scan with only markers would skip OCR.
    if len(text_layer.strip()) < 20 or _readable_ratio(text_layer) < 0.5:
        ocr_text = _ocr_pdf(path)
        if ocr_text:
            return ocr_text, "ocr", None
        return text_layer, "pdf_text", parser_error or "PDF text layer is empty/low quality and OCR returned no usable text"
    return text, "pdf_text", parser_error


def _create_evidence(db, material, evidence_list) -> list:
    created = []
    for ev in evidence_list:
        ev_type = ev.get("type", "TEXTBOOK_CONTENT")
        obj = models.Evidence(
            material_id=material.id,
            evidence_type=ev_type,
            content_reference=ev.get("content", ""),
        )
        db.add(obj)
        created.append(obj)
    # Empty extraction must not be represented by a file name or a fabricated
    # placeholder.  Unit-specific evidence is created in _create_study_units.
    # Assign IDs before they are referenced by StudyUnitEvidence links. With
    # autoflush disabled, IDs stay None until flush; dedupe can skip the later
    # flush entirely, so persist evidence now.
    db.flush()
    return created


def _evidence_type_for_material(material_type: str) -> str:
    return {
        "REAL_EXAM": "EXAM_FREQUENCY",
        "EXAM_SCOPE": "EXAM_SCOPE",
        "TEACHER_PPT": "TEACHER_EMPHASIS",
        "TEXTBOOK": "TEXTBOOK_CONTENT",
        "BILIBILI": "VIDEO_CONTENT",
        "CLASS_NOTES": "CLASS_NOTE",
        "HOMEWORK": "HOMEWORK_PATTERN",
        "LAB": "LAB_CONTENT",
    }.get(material_type, "TEXTBOOK_CONTENT")


def _link_evidence(db, unit_id: str, evidence: models.Evidence, relevance: float = 1.0) -> None:
    existing = db.get(models.StudyUnitEvidence, (unit_id, evidence.id))
    if existing is None:
        db.add(models.StudyUnitEvidence(
            study_unit_id=unit_id, evidence_id=evidence.id, relevance_score=relevance
        ))


def _create_study_units(db, course_id, unit_list, material) -> list:
    units = []
    # Domain grouping: reuse the course's auto domain across multiple materials
    # instead of creating a duplicate per material.
    course = db.get(models.Course, course_id)
    domain = (
        db.query(models.KnowledgeDomain)
        .filter(
            models.KnowledgeDomain.course_id == course_id,
            models.KnowledgeDomain.title == f"{course.name} · 知识点",
        )
        .first()
    )
    if domain is None:
        domain = models.KnowledgeDomain(course_id=course_id, title=f"{course.name} · 知识点")
        db.add(domain)
        db.flush()

    for item in unit_list:
        title = (item.get("title") or "").strip()
        snippet = (item.get("evidence") or "").strip()
        if not title or not snippet:
            continue
        # Incremental add: reuse an existing unit with the same title instead of
        # creating a duplicate, then only attach the new evidence to it.
        existing = (
            db.query(models.StudyUnit)
            .filter(
                models.StudyUnit.course_id == course_id,
                models.StudyUnit.title == title[:200],
            )
            .first()
        )
        if existing is not None:
            unit = existing
        else:
            unit = models.StudyUnit(
                course_id=course_id,
                domain_id=domain.id,
                title=title[:200],
                description=item.get("description"),
                importance_score=0.0,
                difficulty_score=50.0,
                importance_label="LOW",
                mastery_status="UNKNOWN",
                coverage_status="NOT_STARTED",
                estimated_minutes=10,
            )
            db.add(unit)
            db.flush()
        ev = models.Evidence(
            material_id=material.id,
            evidence_type=_evidence_type_for_material(material.type),
            content_reference=snippet[:3000],
            metadata_={"knowledge_title": title[:200], "confidence": item.get("confidence", "UNKNOWN")},
        )
        db.add(ev)
        db.flush()
        _link_evidence(db, unit.id, ev)
        units.append(unit)
    return units


def _create_questions(db, course_id, question_list) -> list:
    questions = []
    for q in question_list:
        content = (q.get("content") or "").strip()
        if not content or len(content) < 8:
            continue
        question_type = q.get("type", "SHORT_ANSWER")
        obj = models.Question(
            course_id=course_id,
            source_type="REAL_EXAM",
            type=question_type,
            content=content,
            options=q.get("options"),
            answer=q.get("answer") or "",
            # Keep the extracted source block internally so the later mapping
            # can produce auditable evidence. It is never shown as a solution.
            explanation=("[SOURCE]" + q["evidence"]) if q.get("evidence") else q.get("explanation"),
            assessment_target="|".join(q.get("knowledge_titles", [])) or q.get("assessment_target"),
        )
        db.add(obj)
        db.flush()
        questions.append(obj)
    return questions


def _map_questions(db, course_id, questions, units, material) -> None:
    # Map questions only when there is an explicit target or a strong lexical
    # match.  The former implementation mapped every question to units[0] when
    # it could not find a match, silently corrupting practice and mastery.
    if not questions:
        return
    all_units = (
        db.query(models.StudyUnit)
        .filter(models.StudyUnit.course_id == course_id)
        .all()
    )
    for q in questions:
        text = (q.content or "") + (q.assessment_target or "")
        targets = q.assessment_target.split("|") if q.assessment_target else []
        ranked = []
        for unit in all_units:
            score = max([_keyword_overlap(text, unit.title)] + [_keyword_overlap(target, unit.title) for target in targets])
            if score >= 0.55:
                ranked.append((unit, score))
        ranked.sort(key=lambda item: -item[1])
        for position, (unit, score) in enumerate(ranked[:3]):
            db.add(models.QuestionStudyUnit(
                question_id=q.id,
                study_unit_id=unit.id,
                is_primary=position == 0,
                relevance_score=score,
            ))
            # The exact question block is first-class evidence for the mapping.
            if q.explanation and q.explanation.startswith("[SOURCE]"):
                snippet = q.explanation.removeprefix("[SOURCE]").strip()
                ev = models.Evidence(
                    material_id=material.id,
                    evidence_type=_evidence_type_for_material(material.type),
                    content_reference=snippet[:3000],
                    metadata_={"knowledge_title": unit.title, "kind": "question_mapping"},
                )
                db.add(ev)
                db.flush()
                _link_evidence(db, unit.id, ev, score)


def _keyword_overlap(text: str, title: str) -> float:
    normalized_text = re.sub(r"\s+", "", (text or "")).lower()
    normalized_title = re.sub(r"\s+", "", (title or "")).lower()
    if len(normalized_title) < 2:
        return 0.0
    if normalized_title in normalized_text:
        return 1.0
    bigrams = {normalized_title[i:i + 2] for i in range(len(normalized_title) - 1)}
    if not bigrams:
        return 0.0
    return len([gram for gram in bigrams if gram in normalized_text]) / len(bigrams)
