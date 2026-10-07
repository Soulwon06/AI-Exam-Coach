import hashlib

from docx import Document

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import models
from app.database import Base
from app.services.course_brain import PARSER_VERSION, parse_material
from app.services.material_service import infer_existing_material_type


def test_parse_material_persists_and_reuses_text_cache(tmp_path):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    path = tmp_path / "notes.txt"
    path.write_text("数据库事务必须满足 ACID 特性。" * 10, encoding="utf-8")
    material = models.Material(
        id="m1", course_id="c1", type="TEXTBOOK", source="upload",
        name="notes.txt", file_path=str(path), status="PENDING",
    )
    session.add(material)
    session.flush()

    first = parse_material(session, material)
    assert first["cached"] is False
    assert first["status"] == "PARSED"
    assert first["method"] == "text"
    assert material.extracted_text.startswith("数据库事务")
    assert material.extracted_char_count == len(material.extracted_text)
    assert material.file_hash == hashlib.sha256(path.read_bytes()).hexdigest()
    assert material.parser_version == PARSER_VERSION

    second = parse_material(session, material)
    assert second["cached"] is True
    assert second["method"] == "text"
    material.parsing_status = "PROCESSING"
    recovered = parse_material(session, material)
    assert recovered["cached"] is True
    assert recovered["status"] == "PARSED"
    session.close()


def test_legacy_doc_is_explicitly_unsupported(tmp_path):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    path = tmp_path / "old.doc"
    path.write_bytes(b"not parsed as docx")
    material = models.Material(
        id="m2", course_id="c1", type="CLASS_NOTES", source="upload",
        name="old.doc", file_path=str(path), status="PENDING",
    )
    session.add(material)
    session.flush()
    result = parse_material(session, material)
    assert result["status"] == "UNSUPPORTED"
    assert result["method"] == "unsupported_doc"
    assert "legacy .doc" in result["error"]
    assert material.extracted_text == ""
    session.close()


def test_legacy_filename_migration_is_conservative():
    assert infer_existing_material_type("2024期末试卷.pdf", "TEXTBOOK") == "REAL_EXAM"
    assert infer_existing_material_type("老师课件.pptx", "TEXTBOOK") == "TEACHER_PPT"
    assert infer_existing_material_type("教材.pdf", "TEXTBOOK") == "TEXTBOOK"
    assert infer_existing_material_type("期末试卷.pdf", "REAL_EXAM") == "REAL_EXAM"
    assert infer_existing_material_type("数据库18.pdf", "TEXTBOOK", "选择题 考试日期") == "REAL_EXAM"
    assert infer_existing_material_type("复习重点(试题).doc", "TEXTBOOK") == "REAL_EXAM"


def test_docx_text_layer_is_parsed_without_ocr(tmp_path):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    path = tmp_path / "review.docx"
    document = Document()
    document.add_paragraph("数据库系统复习重点")
    document.add_paragraph("事务必须满足 ACID 特性。" * 10)
    document.save(path)
    material = models.Material(
        id="m3", course_id="c1", type="REAL_EXAM", source="upload",
        name="review.docx", file_path=str(path), status="PENDING",
    )
    session.add(material)
    session.flush()

    result = parse_material(session, material)

    assert result["status"] == "PARSED"
    assert result["method"] == "docx"
    assert result["cached"] is False
    assert material.parser_method == "docx"
    assert "数据库系统复习重点" in material.extracted_text
    session.close()
